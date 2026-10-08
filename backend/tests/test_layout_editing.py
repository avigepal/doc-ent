import json

import pytest

from app.editing.layout import LayoutResult, LayoutUnsupported, batches, needs_rebuild, reword_batch

docx = pytest.importorskip("docx")
pymupdf = pytest.importorskip("pymupdf")

from docx.oxml.ns import qn  # noqa: E402

from app.editing.docx_layout import edit_docx  # noqa: E402
from app.editing.pdf_layout import edit_pdf  # noqa: E402


class NoModel:
    def chat(self, *args, **kwargs):
        raise AssertionError("a plain replacement must not call the model")


class FakeModel:
    def __init__(self, reply):
        self.reply = reply
        self.calls = []

    def chat(self, system, user, temperature=0.2, **kwargs):
        self.calls.append(user)
        return self.reply if isinstance(self.reply, str) else json.dumps(self.reply)


def final(updates):
    results = [u for u in updates if isinstance(u, LayoutResult)]
    assert results
    return results[-1]


# ---------- which changes keep the layout ----------

@pytest.mark.parametrize("text", ["make it shorter", "use a formal tone", "fix the grammar", "translate it to French"])
def test_rewording_keeps_the_layout(text):
    assert not needs_rebuild(text)


@pytest.mark.parametrize(
    "text", ["add a summary at the top", "remove the second section", "convert it to a table", "make it a bulleted list"]
)
def test_structural_changes_need_a_new_document(text):
    assert needs_rebuild(text)


def test_paragraphs_are_sent_in_batches():
    paragraphs = [(i, "x" * 1500) for i in range(10)]
    sizes = [len(b) for b in batches(paragraphs)]
    assert sum(sizes) == 10 and max(sizes) <= 3


def test_only_ids_that_were_sent_and_real_text_are_taken_from_the_model():
    reply = {"paragraphs": [{"id": 1, "text": "new"}, {"id": 9, "text": "stray"}, {"id": 2, "text": ""}, {"id": "x", "text": "bad"}]}
    assert reword_batch(FakeModel(reply), "do it", [(1, "a"), (2, "b")]) == {1: "new"}


def test_an_unreadable_model_reply_is_given_up_on_cleanly():
    with pytest.raises(LayoutUnsupported):
        reword_batch(FakeModel("not json at all"), "do it", [(1, "a")])


# ---------- Word ----------

def make_docx(path):
    d = docx.Document()
    d.add_heading("About Mac Studio", 0)
    p = d.add_paragraph("The ")
    bold = p.add_run("Mac Studio")
    bold.bold = True
    p.add_run(" is quiet. We like the Mac Studio.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Mac Studio M5"
    table.cell(0, 1).text = "96 GB"
    d.sections[0].header.paragraphs[0].text = "Mac Studio notes"
    d.save(str(path))


def test_docx_replacement_keeps_formatting_and_reaches_tables_and_headers(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    make_docx(src)

    result = final(edit_docx(src, dst, "instead of mac studio i want mac mini", NoModel()))

    assert result.replacements == 5
    out = docx.Document(str(dst))
    assert out.paragraphs[0].text == "About Mac Mini"
    paragraph = out.paragraphs[1]
    assert paragraph.text == "The Mac Mini is quiet. We like the Mac Mini."
    assert [(r.text, r.bold) for r in paragraph.runs][1] == ("Mac Mini", True)  # still bold
    assert out.tables[0].cell(0, 0).text == "Mac Mini M5"
    assert out.sections[0].header.paragraphs[0].text == "Mac Mini notes"


def test_docx_replacing_a_name_that_is_not_there_says_so(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src)
    with pytest.raises(ValueError, match="doesn't appear"):
        list(edit_docx(src, tmp_path / "b.docx", "replace Raspberry Pi with Jetson", NoModel()))


def test_docx_rewording_changes_only_what_the_model_returns(tmp_path):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    make_docx(src)
    # the editor numbers paragraphs by their position in the document
    body = list(docx.Document(str(src)).element.body.iter(qn("w:p")))
    index = next(i for i, p in enumerate(body) if "is quiet" in "".join(t.text or "" for t in p.iter(qn("w:t"))))
    model = FakeModel({"paragraphs": [{"id": index, "text": "The Mac Studio is very quiet. We like the Mac Studio."}]})

    result = final(edit_docx(src, dst, "make it stronger", model))

    assert result.changed == 1
    out = docx.Document(str(dst))
    assert out.paragraphs[1].text == "The Mac Studio is very quiet. We like the Mac Studio."
    assert [(r.text, r.bold) for r in out.paragraphs[1].runs][1] == ("Mac Studio", True)
    assert out.paragraphs[0].text == "About Mac Studio"


def test_docx_where_the_model_changes_nothing_is_reported(tmp_path):
    src = tmp_path / "a.docx"
    make_docx(src)
    with pytest.raises(ValueError, match="nothing"):
        list(edit_docx(src, tmp_path / "b.docx", "make it stronger", FakeModel({"paragraphs": []})))


# ---------- PDF ----------

def make_pdf(path):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "About Mac Studio", fontname="hebo", fontsize=20)
    page.insert_text((72, 140), "The Mac Studio is quiet and we like it.", fontname="helv", fontsize=12)
    page.insert_text((72, 170), "Second line stays exactly as it is.", fontname="helv", fontsize=12)
    page.draw_rect(pymupdf.Rect(60, 200, 300, 260), color=(1, 0, 0), fill=(1, 0.9, 0.9))
    doc.save(str(path))


def pdf_text(path):
    with pymupdf.open(str(path)) as doc:
        return doc[0].get_text("text")


def test_pdf_replacement_changes_every_line_and_keeps_the_page(tmp_path):
    src, dst = tmp_path / "a.pdf", tmp_path / "b.pdf"
    make_pdf(src)

    result = final(edit_pdf(src, dst, "instead of mac studio i want mac mini", NoModel()))

    assert result.replacements == 2
    text = pdf_text(dst)
    assert "About Mac Mini" in text and "The Mac Mini is quiet and we like it." in text
    assert "Studio" not in text
    assert "Second line stays exactly as it is." in text
    with pymupdf.open(str(dst)) as doc:  # the drawing under the text is still there
        assert len(doc[0].get_drawings()) == 1


def test_pdf_name_that_is_not_on_a_line_falls_back_to_the_clean_document_route(tmp_path):
    src = tmp_path / "a.pdf"
    make_pdf(src)
    with pytest.raises(LayoutUnsupported):
        list(edit_pdf(src, tmp_path / "b.pdf", "replace Raspberry Pi with Jetson", NoModel()))


def test_pdf_text_that_needs_shaping_falls_back(tmp_path):
    src = tmp_path / "a.pdf"
    make_pdf(src)
    with pytest.raises(LayoutUnsupported):
        list(edit_pdf(src, tmp_path / "b.pdf", "replace Mac Studio with मैक", NoModel()))


def test_a_pdf_without_text_falls_back(tmp_path):
    doc = pymupdf.open()
    doc.new_page()
    src = tmp_path / "blank.pdf"
    doc.save(str(src))
    with pytest.raises(LayoutUnsupported, match="no text"):
        list(edit_pdf(src, tmp_path / "out.pdf", "make it shorter", FakeModel({"paragraphs": []})))


def test_pdf_rewording_rewrites_only_the_blocks_the_model_returns(tmp_path):
    src, dst = tmp_path / "a.pdf", tmp_path / "b.pdf"
    make_pdf(src)
    # ids follow reading order: 0 the heading, 1 the first body line, 2 the second
    model = FakeModel({"paragraphs": [{"id": 1, "text": "The Mac Studio is quiet."}]})

    result = final(edit_pdf(src, dst, "make it shorter", model))

    assert result.changed == 1
    text = pdf_text(dst)
    assert "The Mac Studio is quiet." in text and "and we like it" not in text
    assert "About Mac Studio" in text and "Second line stays exactly as it is." in text
