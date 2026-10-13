import csv
import io
import zipfile
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main
from app.config import settings
from app.db import get_session
from app.export.tables import (
    NoTablesError,
    Table,
    as_cell_value,
    csv_safe,
    export_tables,
    extract_tables,
    write_csv,
    write_xlsx,
)

ANSWER = """# Summary

Intro text with a | pipe in prose.

## Sites

| Site | Status | Cost |
|:-----|:------:|-----:|
| **North** pump | Operational | 1,200 |
| [South](http://x) pump | Needs repair<br>urgent | 45% |
| Escaped \\| pipe | ok | 007 |

---

## Costs

| a | b |
|---|---|
| 1 | 2 |
"""


# ---------- reading tables out of Markdown ----------


def test_tables_are_found_with_their_headings_and_cleaned_cells():
    sites, costs = extract_tables(ANSWER)

    assert sites.title == "Sites" and costs.title == "Costs"
    assert sites.headers == ["Site", "Status", "Cost"]
    assert sites.rows[0] == ["North pump", "Operational", "1,200"]  # bold removed
    assert sites.rows[1] == ["South pump", "Needs repair\nurgent", "45%"]  # link text kept, <br> a line break
    assert sites.rows[2][0] == "Escaped | pipe"  # an escaped pipe stays in its cell
    assert costs.rows == [["1", "2"]]


def test_prose_with_a_pipe_and_a_rule_is_not_a_table():
    assert extract_tables("Choose a | b\n\n---\n\nno table here") == []
    assert extract_tables("") == []


def test_ragged_rows_are_padded_to_the_widest():
    (table,) = extract_tables("| a | b | c |\n|---|---|---|\n| 1 | 2 |\n| 1 | 2 | 3 | 4 |\n")
    assert len(table.headers) == 4
    assert all(len(row) == 4 for row in table.rows)


def test_a_table_with_no_leading_pipes_is_read():
    (table,) = extract_tables("a | b\n--- | ---\n1 | 2\n")
    assert table.headers == ["a", "b"] and table.rows == [["1", "2"]]


# ---------- numbers ----------


@pytest.mark.parametrize(
    "text, expected",
    [("1,200", 1200), ("12.5", 12.5), ("-3", -3), ("45%", 0.45), ("1.0", 1.0)],
)
def test_plain_numbers_become_numbers(text, expected):
    assert as_cell_value(text) == expected


@pytest.mark.parametrize("text", ["007", "abc", "12345678901234567", "1-2", "A12", "2024-01-05", ""])
def test_ids_text_and_dates_stay_text(text):
    assert as_cell_value(text) == text


# ---------- the workbook ----------


def test_xlsx_has_a_sheet_per_table_a_styled_header_and_real_numbers(tmp_path):
    path = write_xlsx(extract_tables(ANSWER), tmp_path / "out.xlsx")
    workbook = load_workbook(path)

    assert workbook.sheetnames == ["Sites", "Costs"]
    sheet = workbook["Sites"]
    assert [c.value for c in sheet[1]] == ["Site", "Status", "Cost"]
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    assert sheet["C2"].value == 1200 and isinstance(sheet["C2"].value, int)
    assert sheet["C3"].value == 0.45
    assert sheet["C4"].value == "007"  # leading zeros kept
    assert sheet["B3"].value == "Needs repair\nurgent"


def test_sheet_names_are_valid_and_unique(tmp_path):
    tables = [
        Table(headers=["a"], rows=[["1"]], title="Q1: [draft] results/2024"),
        Table(headers=["a"], rows=[["2"]], title="Q1: [draft] results/2024"),
        Table(headers=["a"], rows=[["3"]], title="x" * 80),
        Table(headers=["a"], rows=[["4"]]),
    ]
    names = load_workbook(write_xlsx(tables, tmp_path / "n.xlsx")).sheetnames

    assert len(set(n.lower() for n in names)) == 4
    assert all(len(n) <= 31 and not any(c in n for c in "[]:*?/\\") for n in names)


def test_a_table_with_no_rows_still_gets_a_header(tmp_path):
    path = write_xlsx([Table(headers=["a", "b"], rows=[], title="Empty")], tmp_path / "e.xlsx")
    assert [c.value for c in load_workbook(path)["Empty"][1]] == ["a", "b"]


# ---------- CSV ----------


def test_one_table_is_one_csv_that_excel_reads_with_accents(tmp_path):
    path = write_csv([Table(headers=["Naïve", "n"], rows=[["café, \"quoted\"", "1"]])], tmp_path / "t.csv")

    assert path.suffix == ".csv"
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # a BOM, so Excel gets the encoding right
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows == [["Naïve", "n"], ['café, "quoted"', "1"]]


def test_several_tables_become_a_zip_of_csvs(tmp_path):
    path = write_csv(extract_tables(ANSWER), tmp_path / "answer.csv")

    assert path.suffix == ".zip"
    with zipfile.ZipFile(path) as bundle:
        assert sorted(bundle.namelist()) == ["Costs.csv", "Sites.csv"]
        assert b"North pump" in bundle.read("Sites.csv")
    assert not list(tmp_path.glob("*.csv"))  # no loose files left behind


# ---------- errors ----------


def test_an_answer_without_a_table_cannot_be_exported_as_one(tmp_path):
    with pytest.raises(NoTablesError):
        export_tables("just prose", tmp_path, "answer", "xlsx")
    with pytest.raises(ValueError):
        export_tables(ANSWER, tmp_path, "answer", "pptx")


# ---------- the endpoint ----------


@pytest.fixture
def client(tmp_path, monkeypatch):
    import app.history.store as store

    monkeypatch.setattr(settings, "data_dir", str(tmp_path))
    monkeypatch.setattr(store, "record_export", MagicMock(return_value=7))
    main.app.dependency_overrides[get_session] = lambda: MagicMock()
    test_client = TestClient(main.app)
    test_client.record_export = store.record_export
    yield test_client
    main.app.dependency_overrides.clear()


def _export(client, content, fmt):
    return client.post(
        "/export",
        json={"content": content, "filename": "my answer", "fmt": fmt},
        headers={"Authorization": f"Bearer {settings.bearer_token}"},
    )


def test_the_export_endpoint_returns_an_xlsx(client):
    response = _export(client, ANSWER, "xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "my-answer.xlsx" in response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["Sites", "Costs"]
    assert client.record_export.call_args.kwargs["fmt"] == "xlsx"


def test_several_tables_as_csv_come_back_as_a_zip_recorded_as_one(client):
    response = _export(client, ANSWER, "csv")

    assert response.status_code == 200 and response.headers["content-type"] == "application/zip"
    assert client.record_export.call_args.kwargs["fmt"] == "zip"


def test_one_table_as_csv_is_a_plain_csv(client):
    response = _export(client, "| a | b |\n|---|---|\n| 1 | 2 |\n", "csv")

    assert response.status_code == 200 and response.headers["content-type"].startswith("text/csv")


def test_asking_for_tables_where_there_are_none_is_a_clear_400(client):
    response = _export(client, "No tables in this answer.", "xlsx")

    assert response.status_code == 400
    assert "no table" in response.json()["detail"]


# ---------- text from a document must never run as a formula ----------

DANGEROUS = ["=cmd|' /C calc'!A0", '=HYPERLINK("http://evil","click")', "@SUM(1+1)", "+cmd|x", "-2+3"]


def test_xlsx_keeps_formula_looking_text_as_text(tmp_path):
    table = Table(headers=["x"], rows=[[c] for c in DANGEROUS])
    sheet = load_workbook(write_xlsx([table], tmp_path / "f.xlsx"))["Table 1"]

    for row, text in enumerate(DANGEROUS, start=2):
        assert sheet.cell(row=row, column=1).data_type != "f"
        assert sheet.cell(row=row, column=1).value == text


def test_csv_neutralises_cells_excel_would_run():
    for text in DANGEROUS:
        assert csv_safe(text).startswith("'")
    # ordinary values are left alone, negative numbers included
    for text in ["hello", "-3", "+5", "-1,200.50", "", "a=b"]:
        assert csv_safe(text) == text


def test_written_csv_files_are_neutralised(tmp_path):
    path = write_csv([Table(headers=["x"], rows=[["=1+1"], ["-3"]])], tmp_path / "t.csv")
    rows = list(csv.reader(io.StringIO(path.read_bytes().decode("utf-8-sig"))))
    assert rows == [["x"], ["'=1+1"], ["-3"]]
