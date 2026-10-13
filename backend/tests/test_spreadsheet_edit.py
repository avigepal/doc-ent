"""Editing an attached Excel file from the chat: the whole path through edit_events."""

import json
import sys
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from app.editing.xlsx_ops import OPS_SYSTEM_PROMPT

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class FakeModel:
    """Answers the operations prompt with the next list of operations."""

    def __init__(self, *replies, problem=""):
        self.replies = list(replies)
        self.problem = problem
        self.calls = []

    def chat(self, system, user, temperature=0.2, **kwargs):
        self.calls.append(user)
        assert system == OPS_SYSTEM_PROMPT, f"unexpected prompt: {system[:40]}"
        ops = self.replies.pop(0) if self.replies else []
        return json.dumps({"operations": ops, "problem": self.problem})


def _write_book(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock"
    for row in [["Item", "Qty", "Price"], ["bolt", 10, 2.5], ["nut", 4, 1.25], ["gear", 2, 9]]:
        ws.append(row)
    wb.save(path)


@pytest.fixture
def env(tmp_path, monkeypatch):
    # the editing module imports PyMuPDF (in the Docker image, not in a bare local venv)
    try:
        import pymupdf  # noqa: F401
    except ImportError:
        monkeypatch.setitem(sys.modules, "pymupdf", MagicMock())

    import app.editing.service as service

    raw = tmp_path / "raw"
    raw.mkdir()
    state = SimpleNamespace(exports=[], earlier=[], tmp=tmp_path, raw=raw)
    monkeypatch.setattr(service, "edits_in_conversation", lambda *a, **k: list(state.earlier))

    def fake_record_export(session, **kw):
        record = {**kw, "stored_path": str(kw["stored_path"]), "id": len(state.exports) + 1}
        state.exports.append(SimpleNamespace(**record))
        return len(state.exports)

    monkeypatch.setattr(service, "record_export", fake_record_export)

    def run(name, instruction, llm, mime=XLSX_MIME):
        row = SimpleNamespace(id=1, path=str(raw / name), mime_type=mime)
        session = MagicMock()
        session.execute.return_value.scalars.return_value = [row]
        events = list(service.edit_events(session, instruction=instruction, file_ids=[1], llm=llm, data_dir=str(tmp_path)))
        files = [e["data"] for e in events if e["event"] == "file"]
        message = "".join(e["data"]["text"] for e in events if e["event"] == "token")
        statuses = [e["data"]["text"] for e in events if e["event"] == "status"]
        return files, message, statuses

    state.run = run
    state.service = service
    return state


def test_an_attached_xlsx_is_edited_into_a_new_xlsx_and_the_original_is_untouched(env):
    _write_book(env.raw / "stock.xlsx")
    llm = FakeModel([{"op": "add_column", "header": "Value", "formula": "=B{row}*C{row}"}])

    files, message, _ = env.run("stock.xlsx", "add a value column", llm)

    assert len(files) == 1 and files[0]["fmt"] == "xlsx" and files[0]["name"] == "stock-edited.xlsx"
    out = load_workbook(env.exports[0].stored_path)["Stock"]
    assert out["D1"].value == "Value" and out["D2"].value == "=B2*C2"
    assert load_workbook(env.raw / "stock.xlsx")["Stock"]["D1"].value is None
    assert "Edited **stock.xlsx**" in message and "stock-edited.xlsx" in message
    assert (Path(env.exports[0].stored_path).parent / env.service.LAYOUT_FLAG_NAME).exists()


def test_the_next_edit_in_the_chat_continues_from_the_previous_result(env):
    _write_book(env.raw / "stock.xlsx")
    first = FakeModel([{"op": "add_column", "header": "Value", "formula": "=B{row}*C{row}"}])
    env.run("stock.xlsx", "add a value column", first)
    env.earlier = [env.exports[0]]

    second = FakeModel([{"op": "sort", "column": "Item", "descending": True}])
    files, message, _ = env.run("stock.xlsx", "now sort by item, z to a", second)

    assert files[0]["name"] == "stock-edited-2.xlsx"
    out = load_workbook(env.exports[1].stored_path)["Stock"]
    assert out["D1"].value == "Value"  # the first edit is still there
    assert [c.value for c in out["A"]][1:] == ["nut", "gear", "bolt"]
    assert "continues from your previous edit" in message


def test_from_the_original_starts_over(env):
    _write_book(env.raw / "stock.xlsx")
    env.run("stock.xlsx", "add a value column", FakeModel([{"op": "add_column", "header": "Value", "value": 1}]))
    env.earlier = [env.exports[0]]

    files, message, _ = env.run(
        "stock.xlsx", "from the original, sort by item", FakeModel([{"op": "sort", "column": "Item"}])
    )

    out = load_workbook(env.exports[1].stored_path)["Stock"]
    assert out["D1"].value is None  # not carried over
    assert "continues" not in message


def test_an_old_xls_is_converted_first_and_comes_back_as_xlsx(env, monkeypatch):
    (env.raw / "old.xls").write_bytes(b"legacy bytes")
    converted = []

    def fake_convert(source, target, work_dir, soffice="soffice"):
        converted.append((source.name, target))
        out = work_dir / f"{source.stem}.{target}"
        _write_book(out)
        return out

    monkeypatch.setattr("app.conversion.legacy_office.convert_with_libreoffice", fake_convert)

    files, message, statuses = env.run(
        "old.xls", "uppercase the items", FakeModel([{"op": "transform_column", "column": "Item", "transform": "upper"}]),
        mime="application/vnd.ms-excel",
    )

    assert converted == [("old.xls", "xlsx")]
    assert files[0]["name"] == "old-edited.xlsx"
    assert [c.value for c in load_workbook(env.exports[0].stored_path)["Stock"]["A"]][1:] == ["BOLT", "NUT", "GEAR"]
    assert any("Converting the old .xls" in s for s in statuses)
    assert (env.raw / "old.xls").read_bytes() == b"legacy bytes"


def test_when_the_model_cannot_do_it_the_user_is_told_why_and_no_file_is_made(env):
    _write_book(env.raw / "stock.xlsx")

    files, message, _ = env.run("stock.xlsx", "add each supplier's phone number", FakeModel(problem="there are no suppliers in it"))

    assert files == [] and env.exports == []
    assert "couldn't edit **stock.xlsx**" in message and "no suppliers" in message
    assert not list((env.tmp / "outputs").glob("*"))  # nothing left behind


def test_a_workbook_that_cannot_be_opened_is_reported_not_crashed(env):
    (env.raw / "broken.xlsx").write_bytes(b"this is not a zip file")

    files, message, _ = env.run("broken.xlsx", "sort it", FakeModel())

    assert files == [] and "couldn't edit **broken.xlsx**" in message and "couldn't be opened" in message


def test_what_could_not_be_done_is_said_alongside_what_was(env):
    _write_book(env.raw / "stock.xlsx")
    llm = FakeModel(
        [
            {"op": "set_cell", "cell": "A1", "value": "Product"},
            {"op": "set_cell", "cell": "E2", "value": "=HYPERLINK(\"http://evil.example\",\"x\")"},
        ]
    )

    files, message, _ = env.run("stock.xlsx", "rename and link", llm)

    assert len(files) == 1
    assert "Not done" in message and "not allowed" in message
    out = load_workbook(env.exports[0].stored_path)["Stock"]
    assert out["A1"].value == "Product" and out["E2"].value is None
