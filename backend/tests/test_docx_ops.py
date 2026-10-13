"""Exact edits to a Word file: add, remove or change named parts, nothing else."""

import json

import pytest

docx = pytest.importorskip("docx")

from docx.shared import Pt  # noqa: E402

from app.editing.docx_ops import (  # noqa: E402
    OPS_SYSTEM_PROMPT,
    edit_docx_exact,
    parse_ops,
    read_blocks,
    wants_exact_edit,
)
from app.editing.layout import LayoutResult, LayoutUnsupported  # noqa: E402


class FakeModel:
    def __init__(self, ops, problem=None):
        self.reply = {"operations": ops}
        if problem:
            self.reply["problem"] = problem
        self.prompts = []

    def chat(self, system, user, temperature=0.2, **kwargs):
        assert system == OPS_SYSTEM_PROMPT
        self.prompts.append(user)
        return json.dumps(self.reply)


def make(path):
    d = docx.Document()
    d.add_heading("Quotation", 1)  # [1]
    p = d.add_paragraph("The ")  # [2]
    bold = p.add_run("Mac Studio")
    bold.bold = True
    p.add_run(" is quiet.")
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)
    d.add_heading("Items", 1)  # [3]
    d.add_paragraph("Quiet and small", style="List Bullet")  # [4]
    t = d.add_table(rows=2, cols=2)  # [5]
    t.cell(0, 0).text = "Model"
    t.cell(0, 1).text = "Price"
    t.cell(1, 0).text = "Mac Studio"
    t.cell(1, 1).text = "6,29,900"
    d.add_paragraph("Thank you.")  # [6]
    d.save(str(path))


def run(tmp_path, ops, **kw):
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    make(src)
    results = [u for u in edit_docx_exact(src, dst, "do it", FakeModel(ops, **kw)) if isinstance(u, LayoutResult)]
    return docx.Document(str(dst)), results[-1]


def body_texts(document):
    return [p.text for p in document.paragraphs]


# ---------- reading the document ----------

def test_the_document_is_numbered_in_reading_order_with_tables_as_one_block(tmp_path):
    make(tmp_path / "a.docx")

    blocks = read_blocks(docx.Document(str(tmp_path / "a.docx")))

    assert [(b.id, b.kind) for b in blocks] == [
        (1, "heading 1"), (2, "paragraph"), (3, "heading 1"), (4, "list item"), (5, "table"), (6, "paragraph"),
    ]
    assert blocks[4].text.splitlines()[0] == "    r0: Model | Price"
    assert blocks[1].text == "The Mac Studio is quiet."


# ---------- choosing this engine ----------

@pytest.mark.parametrize(
    "instruction,structural,expected",
    [
        ("add a summary after the introduction", True, True),
        ("shorten paragraph 3", False, True),
        ("change the heading Items to Products", False, True),
        ("translate the whole document to French", False, False),
        ("translate the conclusion to French", False, False),
        ("make it more formal", False, False),
    ],
)
def test_requests_that_name_a_place_or_change_structure_use_exact_edits(instruction, structural, expected):
    assert wants_exact_edit(instruction, structural) is expected


# ---------- reading the model's reply ----------

def test_operations_that_name_a_missing_block_or_are_malformed_are_dropped_with_a_reason():
    data = {
        "operations": [
            {"op": "replace", "id": 2, "text": "ok"},
            {"op": "replace", "id": 99, "text": "x"},
            {"op": "delete", "id": "3"},
            {"op": "explode", "id": 1},
            {"op": "insert_after", "id": 1, "text": "  "},
            {"op": "insert_table_after", "id": 1, "rows": []},
            {"op": "set_cell", "id": 5, "row": -1, "col": 0, "text": "x"},
            "junk",
        ]
    }

    ops, dropped = parse_ops(data, {1, 2, 3, 4, 5, 6})

    assert [(o.kind, o.id) for o in ops] == [("replace", 2)]
    assert len(dropped) == 6 and "there is no block 99" in dropped


# ---------- applying them ----------

def test_a_new_paragraph_is_added_after_the_named_block_and_looks_like_its_neighbours(tmp_path):
    doc, result = run(tmp_path, [{"op": "insert_after", "id": 2, "style": "paragraph", "text": "It also runs cool."}])

    assert body_texts(doc)[:4] == ["Quotation", "The Mac Studio is quiet.", "It also runs cool.", "Items"]
    new_run = doc.paragraphs[2].runs[0]
    assert new_run.font.name == "Arial" and new_run.font.size == Pt(11)  # copied from the paragraph beside it
    assert result.summary == "added 1 new part" and result.changes[0].op == "inserted"


def test_several_insertions_after_the_same_block_keep_their_order(tmp_path):
    doc, _ = run(
        tmp_path,
        [
            {"op": "insert_after", "id": 2, "style": "paragraph", "text": "First."},
            {"op": "insert_after", "id": 2, "style": "paragraph", "text": "Second."},
        ],
    )

    assert body_texts(doc)[1:4] == ["The Mac Studio is quiet.", "First.", "Second."]


def test_a_new_heading_and_bullets_take_the_documents_own_styles(tmp_path):
    doc, _ = run(
        tmp_path,
        [
            {"op": "insert_after", "id": 6, "style": "heading", "level": 1, "text": "Summary"},
            {"op": "insert_after", "id": 6, "style": "bullet", "text": "Quiet"},
        ],
    )

    last = doc.paragraphs[-2:]
    assert (last[0].text, last[0].style.name) == ("Summary", "Heading 1")
    assert (last[1].text, last[1].style.name) == ("Quiet", "List Bullet")


def test_a_paragraph_can_be_removed_and_nothing_else_moves(tmp_path):
    doc, result = run(tmp_path, [{"op": "delete", "id": 6}])

    assert "Thank you." not in body_texts(doc)
    assert body_texts(doc) == ["Quotation", "The Mac Studio is quiet.", "Items", "Quiet and small"]
    assert result.summary == "removed 1 part"


def test_replacing_a_paragraph_keeps_the_formatting_of_the_words_it_keeps(tmp_path):
    doc, _ = run(tmp_path, [{"op": "replace", "id": 2, "text": "The Mac Studio is very quiet."}])

    runs = doc.paragraphs[1].runs
    assert doc.paragraphs[1].text == "The Mac Studio is very quiet."
    assert runs[1].text == "Mac Studio" and runs[1].bold is True  # untouched words keep their own bold


def test_a_table_cell_is_changed_alone(tmp_path):
    doc, result = run(tmp_path, [{"op": "set_cell", "id": 5, "row": 1, "col": 1, "text": "5,99,900"}])

    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert cells == ["Model", "Price", "Mac Studio", "5,99,900"]
    assert result.changes[0].label == "table 5, row 2, column 2"


def test_a_new_table_goes_after_the_named_block_in_the_documents_table_style(tmp_path):
    doc, _ = run(
        tmp_path,
        [{"op": "insert_table_after", "id": 6, "rows": [["Item", "Qty"], ["Mac", "2"]], "header": True}],
    )

    assert len(doc.tables) == 2
    body = list(doc.element.body.iterchildren())
    new_table = doc.tables[1]._tbl
    assert body.index(new_table) > body.index(doc.paragraphs[-1]._p)  # after "Thank you."
    assert [c.text for c in doc.tables[1].rows[0].cells] == ["Item", "Qty"]
    assert doc.tables[1].rows[0].cells[0].paragraphs[0].runs[0].bold is True


def test_what_cannot_be_done_is_reported_and_the_rest_is_still_done(tmp_path):
    doc, result = run(
        tmp_path,
        [
            {"op": "delete", "id": 6},
            {"op": "replace", "id": 5, "text": "a table is not a paragraph"},
            {"op": "replace", "id": 42, "text": "nothing here"},
        ],
    )

    assert "Thank you." not in body_texts(doc)
    assert len(result.changes) == 1
    assert any("table" in s for s in result.skipped) and any("block 42" in s for s in result.skipped)


def test_a_paragraph_with_a_picture_is_not_deleted(tmp_path):
    src, dst = tmp_path / "p.docx", tmp_path / "q.docx"
    d = docx.Document()
    d.add_paragraph("Logo goes here")
    pic = d.add_paragraph("Caption")
    pic.add_run().add_picture(str(_png(tmp_path)))
    d.add_paragraph("End")
    d.save(str(src))

    results = [
        u for u in edit_docx_exact(src, dst, "x", FakeModel([{"op": "delete", "id": 2}, {"op": "delete", "id": 1}]))
        if isinstance(u, LayoutResult)
    ]

    assert [p.text for p in docx.Document(str(dst)).paragraphs] == ["Caption", "End"]
    assert any("picture" in s for s in results[-1].skipped)


def _png(tmp_path):
    import struct
    import zlib

    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = zlib.compress(b"\x00\xff\x00\x00")
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", raw) + chunk(b"IEND", b"")
    path = tmp_path / "dot.png"
    path.write_bytes(png)
    return path


# ---------- when it cannot work ----------

def test_a_reply_with_no_usable_operations_says_why(tmp_path):
    src = tmp_path / "a.docx"
    make(src)

    with pytest.raises(ValueError, match="there is no paragraph 12"):
        list(edit_docx_exact(src, tmp_path / "b.docx", "x", FakeModel([], problem="there is no paragraph 12")))


def test_a_document_with_no_text_is_not_edited_this_way(tmp_path):
    src = tmp_path / "e.docx"
    docx.Document().save(str(src))

    with pytest.raises(LayoutUnsupported):
        list(edit_docx_exact(src, tmp_path / "b.docx", "x", FakeModel([])))


def test_the_model_sees_every_block_numbered(tmp_path):
    src = tmp_path / "a.docx"
    make(src)
    model = FakeModel([{"op": "delete", "id": 6}])

    list(edit_docx_exact(src, tmp_path / "b.docx", "remove the thank you line", model))

    prompt = model.prompts[0]
    assert "[1] heading 1: Quotation" in prompt and "[5] table 3x2" not in prompt and "[5] table 2x2:" in prompt
    assert "[6] paragraph: Thank you." in prompt


# ---------- reviewing the changes ----------

def _edited(tmp_path, ops, style_ops=None):
    from app.editing.changes import save_changes

    src, dst = tmp_path / "orig" / "a.docx", tmp_path / "out" / "a-edited.docx"
    src.parent.mkdir()
    dst.parent.mkdir()
    make(src)
    result = [u for u in edit_docx_exact(src, dst, "do it", FakeModel(ops)) if isinstance(u, LayoutResult)][-1]
    save_changes(dst, src, list(result.ops), list(result.changes), style_ops)
    return src, dst


def test_a_changed_file_lists_its_changes_all_accepted(tmp_path):
    from app.editing.changes import describe_changes

    _, dst = _edited(
        tmp_path,
        [{"op": "insert_after", "id": 2, "style": "paragraph", "text": "New."}, {"op": "delete", "id": 6}],
    )

    changes = describe_changes(dst)

    assert [(c["op"], c["accepted"]) for c in changes] == [("inserted", True), ("deleted", True)]
    assert changes[0]["after"] == "New." and changes[1]["before"] == "Thank you."


def test_rejecting_a_change_rebuilds_the_file_without_it_and_keeps_the_others(tmp_path):
    from app.editing.changes import apply_accepted, describe_changes

    _, dst = _edited(
        tmp_path,
        [{"op": "insert_after", "id": 2, "style": "paragraph", "text": "New."}, {"op": "delete", "id": 6}],
    )

    apply_accepted(dst, [0])  # keep the insertion, reject the deletion

    assert body_texts(docx.Document(str(dst))) == [
        "Quotation", "The Mac Studio is quiet.", "New.", "Items", "Quiet and small", "Thank you.",
    ]
    assert [c["accepted"] for c in describe_changes(dst)] == [True, False]

    apply_accepted(dst, [0, 1])  # and they can be taken back

    assert "Thank you." not in body_texts(docx.Document(str(dst)))


def test_rejecting_everything_gives_back_the_original_text(tmp_path):
    from app.editing.changes import apply_accepted

    src, dst = _edited(tmp_path, [{"op": "replace", "id": 2, "text": "The Mac Studio is very quiet."}])

    apply_accepted(dst, [])

    assert body_texts(docx.Document(str(dst))) == body_texts(docx.Document(str(src)))


def test_style_changes_made_with_the_edit_are_kept_when_changes_are_reviewed(tmp_path):
    from app.editing.changes import apply_accepted
    from app.editing.styling import StyleOp

    _, dst = _edited(
        tmp_path,
        [{"op": "insert_after", "id": 2, "style": "paragraph", "text": "New."}],
        style_ops=[StyleOp("headings", None, {"color": "1F3864", "bold": True})],
    )

    apply_accepted(dst, [])

    heading = docx.Document(str(dst)).paragraphs[0].runs[0]
    assert str(heading.font.color.rgb) == "1F3864"


def test_a_file_whose_original_is_gone_cannot_be_reviewed(tmp_path):
    from app.editing.changes import ChangesUnavailable, apply_accepted

    src, dst = _edited(tmp_path, [{"op": "delete", "id": 6}])
    src.unlink()

    with pytest.raises(ChangesUnavailable):
        apply_accepted(dst, [])
