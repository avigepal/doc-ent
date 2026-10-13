"""Building a spreadsheet from documents: the plan, reading rows out of replies, the workbook, and the
whole run with a fake model."""

import json
import sys
import threading
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main
from app.config import settings
from app.db import get_session
from app.export.tables import Table
from app.search import report, spreadsheet
from app.search.report import ReportDocument, ReportLimits
from app.search.spreadsheet import (
    Column,
    DocumentRows,
    SheetPlan,
    build_table,
    merge_document_rows,
    parse_plan,
    parse_rows,
    preview_markdown,
    sheet_scope_problem,
    spreadsheet_events,
)

COLUMNS = [Column("Invoice no"), Column("Date"), Column("Total")]
PLAN = {"sheet": "Invoices", "one_row_per": "item", "columns": [{"name": c.name} for c in COLUMNS]}
LIMITS = ReportLimits(batch_chars=5000, max_files=50, max_chars=1_000_000, notes_chars=40_000, parallel=2)


# ---------- the plan ----------


def test_a_plan_keeps_the_columns_in_order_with_clean_unique_names():
    plan = parse_plan(
        {
            "sheet": "  Stock   list ",
            "one_row_per": "document",
            "columns": ["Name", {"name": " Qty ", "hint": "how many"}, "name", {"name": ""}, 5, {"nope": 1}, "Source"],
        }
    )
    assert [c.name for c in plan.columns] == ["Name", "Qty"]  # the duplicate, blank, junk and the Source column are gone
    assert plan.columns[1].hint == "how many"
    assert plan.sheet == "Stock list" and plan.one_row_per == "document"


def test_a_plan_defaults_to_one_row_per_item_and_is_capped_at_twenty_columns():
    plan = parse_plan({"columns": [f"c{i}" for i in range(40)]})
    assert plan.one_row_per == "item" and plan.sheet == "Data" and len(plan.columns) == spreadsheet.MAX_COLUMNS


def test_no_columns_means_no_plan():
    assert parse_plan({"columns": []}) is None
    assert parse_plan({}) is None
    assert parse_plan({"columns": "name, date"}) is None


# ---------- rows out of a reply ----------


def test_rows_are_matched_to_columns_ignoring_case_and_spacing_and_empty_rows_are_dropped():
    rows = parse_rows(
        {
            "rows": [
                {"invoice no": "A-1", "DATE": "2024-01-05", "Total": 1200.5, "Other": "ignored"},
                {"Invoice  no": "A-2", "Total": 40.0, "Date": None},
                {"Invoice no": "", "Date": " ", "Total": None},
                "not a row",
                {"Invoice no": ["x", "y"], "Total": True},
            ]
        },
        COLUMNS,
    )
    assert rows[0] == {"Invoice no": "A-1", "Date": "2024-01-05", "Total": "1200.5"}
    assert rows[1] == {"Invoice no": "A-2", "Date": "", "Total": "40"}
    assert rows[2] == {"Invoice no": "x; y", "Date": "", "Total": "yes"}
    assert len(rows) == 3


def test_a_reply_without_rows_gives_none():
    assert parse_rows({}, COLUMNS) == [] and parse_rows({"rows": "x"}, COLUMNS) == []


def test_the_pieces_of_a_long_document_become_one_row_when_a_row_is_a_document():
    rows = [
        {"Invoice no": "A-1", "Date": "", "Total": ""},
        {"Invoice no": "", "Date": "2024-01-05", "Total": "10"},
        {"Invoice no": "B-9", "Date": "2025-01-01", "Total": "99"},  # a later value never overwrites
    ]
    assert merge_document_rows(rows, COLUMNS) == [{"Invoice no": "A-1", "Date": "2024-01-05", "Total": "10"}]
    assert merge_document_rows([], COLUMNS) == []


# ---------- the table ----------


def test_the_table_puts_the_source_first_and_is_cut_at_the_row_limit(monkeypatch):
    plan = SheetPlan("Invoices", "item", COLUMNS)
    results = [
        DocumentRows("a.pdf", [{"Invoice no": "1", "Date": "", "Total": "5"}] * 3),
        DocumentRows("b.pdf", [{"Invoice no": "2", "Date": "", "Total": "6"}] * 3),
    ]
    table, truncated = build_table(plan, results)
    assert table.headers == ["Source", "Invoice no", "Date", "Total"] and len(table.rows) == 6 and not truncated
    assert table.rows[0] == ["a.pdf", "1", "", "5"]

    monkeypatch.setattr(spreadsheet, "MAX_ROWS", 4)
    table, truncated = build_table(plan, results)
    assert len(table.rows) == 4 and truncated


def test_the_preview_is_a_markdown_table_that_survives_pipes_and_line_breaks():
    table = Table(headers=["Source", "Note"], rows=[["a.pdf", "x | y\nz"]])
    assert preview_markdown(table) == "| Source | Note |\n| --- | --- |\n| a.pdf | x \\| y z |"


# ---------- the whole run ----------


class FakeLLM:
    """Plans the columns, then answers each document piece from `by_document`."""

    def __init__(self, by_document, plan=PLAN, fail_for=()):
        self.by_document, self.plan, self.fail_for = by_document, plan, set(fail_for)
        self.lock = threading.Lock()
        self.extract_calls = []

    def chat(self, system, user, temperature=0.2, **kwargs):
        if system == spreadsheet.PLAN_SYSTEM_PROMPT:
            return json.dumps(self.plan)
        assert system == spreadsheet.EXTRACT_SYSTEM_PROMPT
        name = user.split("Document: ", 1)[1].split("\n", 1)[0]
        with self.lock:
            self.extract_calls.append((name, user))
        if name in self.fail_for:
            raise RuntimeError("model unavailable")
        return json.dumps({"rows": self.by_document.get(name, [])})


def _doc(name, *texts):
    return ReportDocument(f"/data/raw/finance/{name}", [("", t) for t in texts], file_id=1)


def _run(request, documents, llm, tmp_path, limits=LIMITS):
    published = []

    def publish(path, source):
        published.append((path, source))
        return {"id": 7, "name": path.name, "fmt": "xlsx", "size_bytes": path.stat().st_size, "source": source,
                "download_path": "/history/exports/7/download"}

    events = list(spreadsheet_events(request, documents, llm, limits, tmp_path / "out", publish))
    files = [e["data"] for e in events if e["event"] == "file"]
    text = "".join(e["data"]["text"] for e in events if e["event"] == "token")
    return events, files, text, published


def test_a_spreadsheet_is_built_from_the_rows_found_in_each_document(tmp_path):
    llm = FakeLLM(
        {
            "inv1.pdf": [
                {"Invoice no": "A-1", "Date": "2024-01-05", "Total": "1200.50"},
                {"Invoice no": "A-2", "Date": "2024-02-01", "Total": "40"},
            ],
            "inv2.pdf": [{"Invoice no": "B-7", "Date": "2024-03-09", "Total": "99"}],
        }
    )
    docs = [_doc("inv1.pdf", "text one"), _doc("inv2.pdf", "text two"), _doc("memo.pdf", "no invoices here")]

    events, files, text, published = _run("every invoice with number, date and total", docs, llm, tmp_path)

    assert events[0]["event"] == "meta" and events[0]["data"]["sources"] == [d.path for d in docs]
    assert events[-1]["event"] == "extra"
    assert [e["event"] for e in events].count("meta") == 1  # the sources are not overwritten
    assert len(files) == 1 and files[0]["fmt"] == "xlsx" and files[0]["name"] == "invoices.xlsx"

    sheet = load_workbook(published[0][0])["Invoices"]
    assert [c.value for c in sheet[1]] == ["Source", "Invoice no", "Date", "Total"]
    assert [c.value for c in sheet[2]] == ["inv1.pdf", "A-1", "2024-01-05", 1200.5]  # a real number
    assert sheet.max_row == 4 and sheet["A4"].value == "inv2.pdf"

    assert "I built **invoices.xlsx** from 2 of 3 documents: 3 rows and 4 columns" in text
    assert "| inv1.pdf | A-1 |" in text  # the preview
    assert "Nothing found in: memo.pdf" in text
    assert "left empty" in text


def test_one_row_per_document_merges_the_pieces_of_a_long_document(tmp_path):
    plan = {"sheet": "People", "one_row_per": "document", "columns": [{"name": "Name"}, {"name": "Email"}]}

    class Pieces(FakeLLM):
        def chat(self, system, user, temperature=0.2, **kwargs):
            if system == spreadsheet.PLAN_SYSTEM_PROMPT:
                return json.dumps(plan)
            return json.dumps({"rows": [{"Name": "Ada"}] if "first part" in user else [{"Email": "ada@example.com"}]})

    small = ReportLimits(batch_chars=40, max_files=50, max_chars=10**6, notes_chars=40_000, parallel=2)
    doc = ReportDocument("/data/raw/cv.pdf", [("", "first part " + "x" * 30), ("", "second part " + "y" * 30)])

    _, files, _, published = _run("name and email of each person", [doc], Pieces({}), tmp_path, limits=small)

    sheet = load_workbook(published[0][0])["People"]
    assert sheet.max_row == 2 and [c.value for c in sheet[2]] == ["cv.pdf", "Ada", "ada@example.com"]


def test_text_from_a_document_that_looks_like_a_formula_stays_text(tmp_path):
    llm = FakeLLM({"inv1.pdf": [{"Invoice no": "=cmd|' /C calc'!A0", "Date": "", "Total": '=HYPERLINK("http://evil","x")'}]})
    _, _, _, published = _run("invoices", [_doc("inv1.pdf", "t")], llm, tmp_path)

    sheet = load_workbook(published[0][0])["Invoices"]
    assert sheet["B2"].data_type != "f" and sheet["B2"].value == "=cmd|' /C calc'!A0"
    assert sheet["D2"].data_type != "f"


def test_a_part_that_cannot_be_read_is_reported_and_the_rest_is_still_built(tmp_path):
    llm = FakeLLM({"ok.pdf": [{"Invoice no": "A-1", "Date": "", "Total": "5"}]}, fail_for=["bad.pdf"])
    _, files, text, _ = _run("invoices", [_doc("ok.pdf", "t"), _doc("bad.pdf", "t")], llm, tmp_path)

    assert len(files) == 1
    assert "Some parts couldn't be read (bad.pdf)" in text
    assert "Nothing found in: bad.pdf" not in text  # unreadable is not the same as empty


def test_nothing_found_makes_no_file(tmp_path):
    events, files, text, published = _run("invoices", [_doc("memo.pdf", "t")], FakeLLM({}), tmp_path)

    assert files == [] and published == []
    assert "found nothing that fits" in text
    assert not (tmp_path / "out").exists()


def test_a_request_that_names_nothing_to_collect_asks_what_the_columns_are(tmp_path):
    events, files, text, _ = _run("do something", [_doc("a.pdf", "t")], FakeLLM({}, plan={"columns": []}), tmp_path)
    assert files == [] and "what the spreadsheet should contain" in text


def test_an_unreadable_plan_is_explained_not_raised(tmp_path):
    class Broken(FakeLLM):
        def chat(self, system, user, temperature=0.2, **kwargs):
            return "not json at all"

    _, files, text, _ = _run("invoices", [_doc("a.pdf", "t")], Broken({}), tmp_path)
    assert files == [] and "couldn't work out the columns" in text


def test_every_piece_is_sent_with_the_columns_and_the_request(tmp_path):
    llm = FakeLLM({})
    _run("every invoice with number, date and total", [_doc("a.pdf", "invoice text")], llm, tmp_path)

    _, user = llm.extract_calls[0]
    assert "Invoice no" in user and "Date" in user and "Total" in user
    assert "one per item" in user and "every invoice with number, date and total" in user and "invoice text" in user


# ---------- the size of the job ----------


def test_the_scope_is_checked_before_anything_is_read():
    assert "couldn't find" in sheet_scope_problem([], LIMITS).lower() or sheet_scope_problem([], LIMITS)
    many = [_doc(f"{i}.pdf", "t") for i in range(51)]
    assert "51 documents" in sheet_scope_problem(many, LIMITS)
    big = [_doc("a.pdf", "x" * 2_000_000)]
    assert "characters" in sheet_scope_problem(big, LIMITS)
    assert sheet_scope_problem([_doc("a.pdf", "t")], LIMITS) is None


# ---------- through the chat endpoint ----------


@pytest.fixture
def client(tmp_path, monkeypatch):
    try:
        import pymupdf  # noqa: F401
    except ImportError:
        monkeypatch.setitem(sys.modules, "pymupdf", MagicMock())

    import app.history.store as store
    import app.tasks.correlate as correlate

    llm = FakeLLM({"inv1.pdf": [{"Invoice no": "A-1", "Date": "2024-01-05", "Total": "10"}]})
    monkeypatch.setattr(correlate, "_text_llm", llm)
    monkeypatch.setattr(settings, "data_dir", str(tmp_path))
    monkeypatch.setattr(store, "record_query", MagicMock(return_value=42))
    monkeypatch.setattr(store, "link_exports", MagicMock())
    monkeypatch.setattr(store, "record_export", MagicMock(return_value=9))
    monkeypatch.setattr(
        report, "load_report_documents", lambda session, raw_dir, folders, author, title, file_ids: [_doc("inv1.pdf", "t")]
    )
    monkeypatch.setattr(main, "_attachment_names", lambda session, ids: [])
    monkeypatch.setattr(main, "_catalog_answer", MagicMock(side_effect=AssertionError("catalog path taken")))
    main.app.dependency_overrides[get_session] = lambda: MagicMock()
    test_client = TestClient(main.app)
    test_client.store = store
    yield test_client
    main.app.dependency_overrides.clear()


def _stream(client, **payload):
    response = client.post("/query/stream", json=payload, headers={"Authorization": f"Bearer {settings.bearer_token}"})
    assert response.status_code == 200
    events = []
    for block in response.text.strip().split("\n\n"):
        name, data = block.split("\n", 1)
        events.append((name.removeprefix("event: "), json.loads(data.removeprefix("data: "))))
    return events


def test_the_spreadsheet_toggle_builds_a_file_whatever_the_message_looks_like(client):
    events = _stream(client, question="list the files", mode="spreadsheet", folders=["finance"], k=5)
    names = [name for name, _ in events]

    route = next(data for name, data in events if name == "route")
    assert route["action"] == "spreadsheet" and route["source"] == "user"
    file_data = next(data for name, data in events if name == "file")
    assert file_data["fmt"] == "xlsx" and file_data["id"] == 9 and file_data["name"] == "invoices.xlsx"
    assert names[-1] == "done"
    # the file is linked to the question it answers, so it shows up in History
    client.store.link_exports.assert_called_once()
    assert client.store.link_exports.call_args.args[1] == [9]


def test_a_direct_chat_has_no_documents_to_build_a_spreadsheet_from(client):
    events = _stream(client, question="invoices", mode="spreadsheet", chat_only=True, k=5)

    assert not any(name == "file" for name, _ in events)
    text = "".join(data["text"] for name, data in events if name == "token")
    assert text == report.NO_DOCUMENTS_MESSAGE
