"""Exact edits to a Word file: add, remove or change specific parts of it while
everything else stays as it was.

The model does not rewrite the document. It is shown the document as numbered
blocks (paragraphs and tables, in reading order) and answers with a short list
of operations on those numbers:

  replace            the text of paragraph N
  insert_after       a new paragraph / heading / bullet after block N
  insert_table_after a new table after block N
  set_cell           the text of one table cell
  delete             paragraph or table N

Code applies each operation to the Word file's own XML. A new paragraph is a
copy of a similar one already in the document (same style, fonts, spacing,
numbering), so it looks like it was always there; a replaced text keeps the
formatting of the words it replaces (see textspans.py). Whatever the model
does not name is not touched, and every operation that was applied is recorded
as a Change, so the dashboard can show what happened.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.editing.docx_layout import _own_runs, _paragraph_text, _run_text, _set_run_text, _write_spans
from app.editing.engine import EditProgress
from app.editing.layout import LayoutResult, LayoutUnsupported, ParagraphLLM, chat_json
from app.editing.textspans import diff_spans

_P, _TBL, _TR, _TC, _R, _T = qn("w:p"), qn("w:tbl"), qn("w:tr"), qn("w:tc"), qn("w:r"), qn("w:t")
_PPR, _RPR, _NUMPR, _SECTPR = qn("w:pPr"), qn("w:rPr"), qn("w:numPr"), qn("w:sectPr")
_DRAWING = (qn("w:drawing"), qn("w:pict"), qn("w:object"))

MAX_OPERATIONS = 200
MAX_PROMPT_CHARS = 40_000
MAX_BLOCK_CHARS = 1_500  # a longer paragraph is shown cut; replacing it is then refused
MAX_BLOCKS = 600

# ---------- which requests are worth this engine ----------

# a request that names a place in the document, or adds or removes something
_PLACE = re.compile(
    r"\b(?:paragraphs?|para|headings?|sections?|chapters?|tables?|rows?|columns?|cells?|bullets?|bulleted|"
    r"list\s+items?|intro(?:duction)?|conclusion|summary|abstract|title|first|second|third|last|final|"
    r"beginning|end\s+of|top\s+of|bottom\s+of)\b",
    re.IGNORECASE,
)
# a change of every word in the document: handled paragraph by paragraph, tables and headers included
_EVERYWHERE = re.compile(
    r"\b(?:translate|proofread|spell[\s-]?check|fix\s+(?:all\s+)?(?:the\s+)?(?:grammar|spelling|typos)|"
    r"throughout|entire|whole\s+(?:document|file)|everywhere)\b",
    re.IGNORECASE,
)


def wants_exact_edit(instruction: str, needs_structure: bool) -> bool:
    """True when this request is better done as exact operations on named parts of the
    document than as a rewording of every paragraph."""
    if needs_structure:
        return True
    return bool(_PLACE.search(instruction)) and not _EVERYWHERE.search(instruction)


# ---------- reading the document as numbered blocks ----------

@dataclass
class Block:
    id: int
    kind: str  # "heading 1", "paragraph", "list item", "table"
    text: str  # paragraphs; for a table, its rows as text
    element: object
    rows: int = 0
    cols: int = 0


def _style_name(doc, element) -> str:
    try:
        return (Paragraph(element, doc._body).style.name or "").strip()
    except Exception:
        return ""


def _is_list_item(doc, element) -> bool:
    ppr = element.find(_PPR)
    if ppr is not None and ppr.find(_NUMPR) is not None:
        return True
    return "list" in _style_name(doc, element).lower()


def _heading_level(doc, element) -> int | None:
    name = _style_name(doc, element)
    if name.lower() == "title":
        return 0
    match = re.match(r"heading\s*(\d)$", name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _table_text(table_el) -> tuple[str, int, int]:
    rows = []
    cols = 0
    for tr in table_el.iterchildren(_TR):
        cells = ["".join(_paragraph_text(p) for p in tc.iter(_P)).strip() for tc in tr.iterchildren(_TC)]
        cols = max(cols, len(cells))
        rows.append(cells)
    return "\n".join(f"    r{i}: " + " | ".join(r) for i, r in enumerate(rows)), len(rows), cols


def read_blocks(doc) -> list[Block]:
    """The document's top-level paragraphs and tables, in order, numbered from 1.
    Blank paragraphs are not numbered (nothing to refer to)."""
    blocks: list[Block] = []
    for element in doc.element.body.iterchildren():
        if element.tag == _P:
            text = _paragraph_text(element)
            if not text.strip():
                continue
            level = _heading_level(doc, element)
            if level is not None:
                kind = "title" if level == 0 else f"heading {level}"
            elif _is_list_item(doc, element):
                kind = "list item"
            else:
                kind = "paragraph"
            blocks.append(Block(len(blocks) + 1, kind, text, element))
        elif element.tag == _TBL:
            text, rows, cols = _table_text(element)
            blocks.append(Block(len(blocks) + 1, "table", text, element, rows, cols))
    return blocks


def blocks_prompt(blocks: list[Block]) -> str:
    lines = []
    for b in blocks:
        if b.kind == "table":
            lines.append(f"[{b.id}] table {b.rows}x{b.cols}:\n{b.text}")
        else:
            shown = b.text if len(b.text) <= MAX_BLOCK_CHARS else b.text[:MAX_BLOCK_CHARS] + " …(cut)"
            lines.append(f"[{b.id}] {b.kind}: {shown}")
    return "\n".join(lines)


# ---------- operations ----------

@dataclass
class Op:
    kind: str  # replace | insert_after | insert_table_after | set_cell | delete
    id: int
    text: str = ""
    style: str = "paragraph"  # insert_after: paragraph | heading | bullet | numbered
    level: int = 1
    rows: list = field(default_factory=list)
    header: bool = True
    row: int = 0
    col: int = 0


OPS_SYSTEM_PROMPT = (
    "You edit a Word document with exact operations. You get a request and the document as numbered blocks "
    "(paragraphs, headings, list items and tables, in reading order). Reply with JSON only: "
    '{"operations": [ ... ]}\n'
    "Operations (\"id\" is a block number from the list):\n"
    '  {"op": "replace", "id": N, "text": "..."}   the complete new text of paragraph N\n'
    '  {"op": "insert_after", "id": N, "style": "paragraph|heading|bullet|numbered", "level": 1, "text": "..."}   '
    "a new paragraph after block N (level only for headings; one operation per paragraph, in order)\n"
    '  {"op": "insert_table_after", "id": N, "rows": [["a","b"],["c","d"]], "header": true}\n'
    '  {"op": "set_cell", "id": N, "row": 0, "col": 1, "text": "..."}   one cell of table N (counting from 0)\n'
    '  {"op": "delete", "id": N}   remove paragraph or table N\n'
    "Rules:\n"
    "- Do exactly what the request says and nothing else: name only the blocks it concerns. "
    "\"Paragraph 3\" means block [3]; \"the intro\" means the block that is the introduction.\n"
    "- Keep the document's language. Write new text in the document's own voice and do not invent facts; "
    "a summary or section you are asked to add must come from what the document says.\n"
    "- To change part of a paragraph, replace it with its full new text.\n"
    "- A block cut with …(cut) cannot be replaced.\n"
    "- If the request cannot be done with these operations, or names something that is not in the list, "
    'reply {"operations": [], "problem": "one sentence saying why"}.'
)


def parse_ops(data: dict, valid_ids: set[int]) -> tuple[list[Op], list[str]]:
    """The valid operations in the model's reply, and a sentence for each one dropped."""
    ops: list[Op] = []
    dropped: list[str] = []
    for raw in (data.get("operations") or [])[:MAX_OPERATIONS]:
        if not isinstance(raw, dict):
            continue
        kind = str(raw.get("op", "")).strip().lower()
        ident = raw.get("id")
        if kind not in ("replace", "insert_after", "insert_table_after", "set_cell", "delete"):
            dropped.append(f"unknown operation “{raw.get('op')}”")
            continue
        if isinstance(ident, bool) or not isinstance(ident, int) or ident not in valid_ids:
            dropped.append(f"there is no block {ident}")
            continue
        text = raw.get("text")
        text = text if isinstance(text, str) else ""
        if kind in ("replace", "insert_after") and not text.strip():
            dropped.append(f"block {ident}: no text given")
            continue
        op = Op(kind, ident, text)
        if kind == "insert_after":
            style = str(raw.get("style", "paragraph")).lower()
            op.style = style if style in ("paragraph", "heading", "bullet", "numbered") else "paragraph"
            level = raw.get("level")
            op.level = level if isinstance(level, int) and not isinstance(level, bool) and 1 <= level <= 6 else 1
        elif kind == "insert_table_after":
            rows = raw.get("rows")
            if (
                not isinstance(rows, list)
                or not rows
                or not all(isinstance(r, list) and r for r in rows)
                or max(len(r) for r in rows) > 12
                or len(rows) > 60
            ):
                dropped.append(f"block {ident}: the table was not usable")
                continue
            width = max(len(r) for r in rows)
            op.rows = [[str(c) if c is not None else "" for c in r] + [""] * (width - len(r)) for r in rows]
            op.header = raw.get("header") is not False
        elif kind == "set_cell":
            row, col = raw.get("row"), raw.get("col")
            if not all(isinstance(v, int) and not isinstance(v, bool) and v >= 0 for v in (row, col)):
                dropped.append(f"block {ident}: no cell given")
                continue
            op.row, op.col = row, col
        ops.append(op)
    return ops, dropped


# ---------- applying them ----------

@dataclass(frozen=True)
class Change:
    """One thing that was done to the document, for the change list."""

    op: str  # replaced | inserted | deleted | cell
    block: int
    before: str
    after: str
    label: str  # "paragraph 3", "heading after block 2", "table 1, row 2 col 1"
    op_index: int = -1  # which operation made it (changes.py rebuilds the file from the accepted ones)


def _text_run_template(element):
    """A copy of the first run with text in a paragraph (its formatting), emptied."""
    for run in element.iter(_R):
        if run.findall(_T):
            template = copy.deepcopy(run)
            for child in list(template):
                if child.tag != _RPR:
                    template.remove(child)
            return template
    return None


def _new_paragraph_like(template_el, text: str):
    """A paragraph with `template_el`'s properties (style, numbering, spacing) and the first run's look."""
    new = copy.deepcopy(template_el)
    run_template = _text_run_template(template_el)
    for child in list(new):
        if child.tag != _PPR:
            new.remove(child)
    ppr = new.find(_PPR)
    if ppr is not None:  # a copied section break would split the document
        for sect in ppr.findall(_SECTPR):
            ppr.remove(sect)
    run = OxmlElement("w:r") if run_template is None else run_template
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(t)
    new.append(run)
    return new


def _find_template(doc, blocks: list[Block], near: int, style: str, level: int):
    """An existing paragraph that a new one of this kind should copy: the closest of the right kind."""
    paragraphs = [b for b in blocks if b.kind != "table"]
    order = sorted(paragraphs, key=lambda b: abs(b.id - near))

    def pick(test):
        return next((b.element for b in order if test(b)), None)

    if style == "heading":
        found = pick(lambda b: _heading_level(doc, b.element) == level)
        return found if found is not None else pick(lambda b: _heading_level(doc, b.element) is not None)
    if style in ("bullet", "numbered"):
        return pick(lambda b: b.kind == "list item")
    return pick(lambda b: b.kind == "paragraph")


def _style_paragraph(doc, element, style: str, level: int) -> bool:
    """Give `element` a built-in style when the document has it; True when it did."""
    names = {
        "heading": f"Heading {level}",
        "bullet": "List Bullet",
        "numbered": "List Number",
        "paragraph": "Normal",
    }
    try:
        Paragraph(element, doc._body).style = doc.styles[names[style]]
        return True
    except Exception:
        return False


def _remove(element) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _can_delete(element) -> str | None:
    """Why this paragraph/table can't simply be removed, or None."""
    ppr = element.find(_PPR) if element.tag == _P else None
    if ppr is not None and ppr.find(_SECTPR) is not None:
        return "it carries a section break"
    if any(True for tag in _DRAWING for _ in element.iter(tag)):
        return "it holds a picture or object"
    return None


def apply_ops(doc, blocks: list[Block], ops: list[Op]) -> tuple[list[Change], list[str]]:
    by_id = {b.id: b for b in blocks}
    changes: list[Change] = []
    skipped: list[str] = []
    gone: set[int] = set()
    last_inserted: dict[int, object] = {}  # block id -> the element last placed after it
    op_index = 0

    def add(change: Change) -> None:
        changes.append(replace(change, op_index=op_index))

    def anchor_for(ident: int):
        return last_inserted.get(ident, by_id[ident].element)

    for op_index, op in enumerate(ops):
        block = by_id[op.id]
        if op.id in gone:
            skipped.append(f"block {op.id} was already removed")
            continue
        if op.kind == "replace":
            if block.kind == "table":
                skipped.append(f"block {op.id} is a table; change one cell at a time")
            elif len(block.text) > MAX_BLOCK_CHARS:
                skipped.append(f"block {op.id} is too long to replace safely")
            else:
                spans = diff_spans(block.text, op.text)
                if spans:
                    _write_spans(block.element, spans)
                    add(Change("replaced", op.id, block.text, op.text, f"{block.kind} {op.id}"))
        elif op.kind == "delete":
            why = _can_delete(block.element)
            if why:
                skipped.append(f"block {op.id} wasn't removed because {why}")
                continue
            _remove(block.element)
            gone.add(op.id)
            add(Change("deleted", op.id, block.text if block.kind != "table" else "(table)", "", f"{block.kind} {op.id}"))
        elif op.kind == "insert_after":
            template = _find_template(doc, blocks, op.id, op.style, op.level)
            if template is None:
                template = block.element if block.element.tag == _P else None
            if template is None:
                skipped.append(f"couldn't find a paragraph to copy the look from for the text after block {op.id}")
                continue
            new = _new_paragraph_like(template, op.text)
            # a heading/list wanted but the nearest paragraph is a different kind: use the built-in style
            wrong_kind = (
                (op.style == "heading" and _heading_level(doc, template) != op.level)
                or (op.style in ("bullet", "numbered") and not _is_list_item(doc, template))
                or (op.style == "paragraph" and (_heading_level(doc, template) is not None or _is_list_item(doc, template)))
            )
            if wrong_kind:
                if op.style == "paragraph":
                    ppr = new.find(_PPR)
                    if ppr is not None and ppr.find(_NUMPR) is not None:
                        ppr.remove(ppr.find(_NUMPR))
                if not _style_paragraph(doc, new, op.style, op.level) and op.style == "heading":
                    for run in new.iter(_R):  # the document has no heading style: at least make it stand out
                        rpr = run.find(_RPR)
                        if rpr is None:
                            rpr = OxmlElement("w:rPr")
                            run.insert(0, rpr)
                        rpr.append(OxmlElement("w:b"))
            anchor_for(op.id).addnext(new)
            last_inserted[op.id] = new
            add(Change("inserted", op.id, "", op.text, f"new {op.style if op.style != 'paragraph' else 'paragraph'} after block {op.id}"))
        elif op.kind == "insert_table_after":
            table = doc.add_table(rows=len(op.rows), cols=len(op.rows[0]))
            existing = next((Table(b.element, doc._body) for b in blocks if b.kind == "table"), None)
            try:
                table.style = existing.style if existing is not None and existing.style is not None else doc.styles["Table Grid"]
            except Exception:
                pass
            for r, row in enumerate(op.rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    if op.header and r == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            anchor_for(op.id).addnext(table._tbl)
            last_inserted[op.id] = table._tbl
            add(Change("inserted", op.id, "", "\n".join(" | ".join(r) for r in op.rows), f"new table after block {op.id}"))
        elif op.kind == "set_cell":
            if block.kind != "table":
                skipped.append(f"block {op.id} is not a table")
                continue
            table = Table(block.element, doc._body)
            try:
                cell = table.cell(op.row, op.col)
            except IndexError:
                skipped.append(f"table {op.id} has no row {op.row}, column {op.col}")
                continue
            before = cell.text
            paragraph = cell.paragraphs[0]._p
            if _own_runs(paragraph):
                spans = diff_spans(_paragraph_text(paragraph), op.text)
                if not spans and len(cell.paragraphs) == 1:
                    continue  # already says that
                if spans:
                    _write_spans(paragraph, spans)
            else:
                cell.paragraphs[0].add_run(op.text)
            for extra in cell.paragraphs[1:]:  # a cell's other paragraphs would keep the old text
                _remove(extra._p)
            add(Change("cell", op.id, before, op.text, f"table {op.id}, row {op.row + 1}, column {op.col + 1}"))
    return changes, skipped


def summarize(changes: list[Change]) -> str:
    counts = {"replaced": 0, "inserted": 0, "deleted": 0, "cell": 0}
    for c in changes:
        counts[c.op] += 1
    parts = []
    if counts["replaced"]:
        parts.append(f"changed {counts['replaced']} paragraph{'s' if counts['replaced'] != 1 else ''}")
    if counts["cell"]:
        parts.append(f"changed {counts['cell']} table cell{'s' if counts['cell'] != 1 else ''}")
    if counts["inserted"]:
        parts.append(f"added {counts['inserted']} new part{'s' if counts['inserted'] != 1 else ''}")
    if counts["deleted"]:
        parts.append(f"removed {counts['deleted']} part{'s' if counts['deleted'] != 1 else ''}")
    return ", ".join(parts)


def edit_docx_exact(src: Path, dst: Path, instruction: str, llm: ParagraphLLM) -> Iterator[EditProgress | LayoutResult]:
    doc = Document(str(src))
    blocks = read_blocks(doc)
    if not blocks:
        raise LayoutUnsupported("it has no text to edit")
    if len(blocks) > MAX_BLOCKS:
        raise LayoutUnsupported(f"it is too long to edit this way ({len(blocks)} blocks)")
    listing = blocks_prompt(blocks)
    if len(listing) > MAX_PROMPT_CHARS:
        raise LayoutUnsupported("it is too long to edit this way")

    yield EditProgress("Working out exactly what to change")
    data = chat_json(llm, OPS_SYSTEM_PROMPT, f"Request:\n{instruction}\n\nDocument:\n{listing}", temperature=0.1)
    ops, dropped = parse_ops(data, {b.id for b in blocks})
    if not ops:
        problem = data.get("problem") if isinstance(data.get("problem"), str) else ""
        reason = problem or (dropped[0] if dropped else "")
        raise ValueError(reason or "I couldn't tell what to change in the file for that request")

    yield EditProgress(f"Making {len(ops)} change{'s' if len(ops) != 1 else ''}")
    changes, skipped = apply_ops(doc, blocks, ops)
    if not changes:
        raise ValueError(skipped[0] if skipped else "nothing in the file needed to change for that request")
    doc.save(str(dst))

    yield LayoutResult(
        changed=len(changes),
        summary=summarize(changes),
        changes=tuple(changes),
        skipped=tuple((dropped + skipped)[:5]),
        ops=tuple(ops),
    )
