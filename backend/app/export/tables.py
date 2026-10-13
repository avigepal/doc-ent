"""Tables in an answer -> a spreadsheet.

An answer is Markdown; its tables are GitHub-style pipe tables. They are pulled out and
written as one .xlsx (a sheet per table, with a styled header row and real numbers) or as
CSV (a single table is one .csv; several become a .zip of one .csv each).

The same workbook writer serves "make a spreadsheet from my documents" (app/spreadsheet),
so a table of rows that did not come from Markdown can be written the same way.
"""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Excel's limits
MAX_SHEET_NAME = 31
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

_SEPARATOR_CELL = re.compile(r"^:?-{1,}:?$")
_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")
_NUMBER = re.compile(r"^[-+]?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^[-+]?\d+(?:\.\d+)?$")
_PERCENT = re.compile(r"^[-+]?\d+(?:\.\d+)?\s?%$")


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]
    title: str = ""
    sheet: str = field(default="", repr=False)


class NoTablesError(ValueError):
    pass


# ---------- reading Markdown tables ----------


def _split_row(line: str) -> list[str]:
    """The cells of one pipe-table line (escaped pipes stay inside their cell)."""
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith("\\|"):
        text = text[:-1]
    cells = re.split(r"(?<!\\)\|", text)
    return [cell.strip().replace("\\|", "|") for cell in cells]


def _is_separator(line: str) -> bool:
    """The `|---|:--:|` line under a table's header."""
    if "|" not in line:
        return False
    cells = [c.replace(" ", "") for c in _split_row(line) if c.strip()]
    return bool(cells) and all(_SEPARATOR_CELL.match(c) for c in cells)


_MARKDOWN_NOISE = (
    (re.compile(r"<br\s*/?>", re.IGNORECASE), "\n"),
    (re.compile(r"\[([^\]]*)\]\([^)]*\)"), r"\1"),  # a link keeps its text
    (re.compile(r"(\*\*|__)(.+?)\1"), r"\2"),  # bold
    (re.compile(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])"), r"\1"),  # italic
    (re.compile(r"`([^`]*)`"), r"\1"),  # code
)


def clean_cell(text: str) -> str:
    for pattern, replacement in _MARKDOWN_NOISE:
        text = pattern.sub(replacement, text)
    return text.strip()


def extract_tables(markdown: str) -> list[Table]:
    """Every pipe table in the Markdown, in order. A table's title is the heading just above it."""
    lines = markdown.splitlines()
    tables: list[Table] = []
    heading = ""
    i = 0
    while i < len(lines):
        line = lines[i]
        found = _HEADING.match(line)
        if found:
            heading = clean_cell(found.group(1))
        is_table_start = "|" in line and i + 1 < len(lines) and _is_separator(lines[i + 1])
        if is_table_start:
            headers = [clean_cell(c) for c in _split_row(line)]
            rows: list[list[str]] = []
            j = i + 2
            while j < len(lines) and "|" in lines[j] and lines[j].strip():
                cells = [clean_cell(c) for c in _split_row(lines[j])]
                cells = (cells + [""] * len(headers))[: max(len(headers), len(cells))]
                rows.append(cells)
                j += 1
            width = max([len(headers)] + [len(r) for r in rows])
            headers = headers + [""] * (width - len(headers))
            rows = [r + [""] * (width - len(r)) for r in rows]
            tables.append(Table(headers=headers, rows=rows, title=heading))
            heading = ""
            i = j
            continue
        i += 1
    return tables


# ---------- writing ----------


def _sheet_names(tables: list[Table]) -> list[str]:
    """One valid, unique sheet name per table (Excel: at most 31 characters, none of []:*?/\\)."""
    used: set[str] = set()
    names = []
    for n, table in enumerate(tables, start=1):
        base = _BAD_SHEET_CHARS.sub(" ", table.title or table.sheet or f"Table {n}").strip() or f"Table {n}"
        base = base[:MAX_SHEET_NAME].strip()
        name, k = base, 2
        while name.lower() in used:
            suffix = f" ({k})"
            name = base[: MAX_SHEET_NAME - len(suffix)].rstrip() + suffix
            k += 1
        used.add(name.lower())
        names.append(name)
    return names


def as_cell_value(text: str):
    """A number for text that is plainly a number ("1,200", "12.5", "45%"), the text otherwise.
    Leading zeros ("007"), long ids and anything with letters stay text."""
    value = text.strip()
    if not value:
        return ""
    if _PERCENT.match(value):
        return float(value.rstrip("% ").strip()) / 100
    if _NUMBER.match(value):
        digits = value.replace(",", "").lstrip("+-")
        whole = digits.split(".")[0]
        if (len(whole) > 1 and whole.startswith("0")) or len(whole) > 15:
            return value
        number = float(value.replace(",", ""))
        return int(number) if number.is_integer() and "." not in value else number
    return value


_HEADER_FILL = PatternFill("solid", start_color="E8ECF4")


def write_xlsx(tables: list[Table], path: Path) -> Path:
    """One sheet per table: a bold, shaded header row that stays in view, real numbers, readable widths."""
    if not tables:
        raise NoTablesError("there is no table to put in a spreadsheet")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for table, name in zip(tables, _sheet_names(tables)):
        sheet = workbook.create_sheet(name)
        for r, values in enumerate([table.headers] + [[as_cell_value(c) for c in row] for row in table.rows], start=1):
            for c, value in enumerate(values, start=1):
                cell = sheet.cell(row=r, column=c, value=value)
                if isinstance(value, str) and value.startswith("="):
                    cell.data_type = "s"  # text from a document must never become a live formula
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for index, header in enumerate(table.headers, start=1):
            longest = max([len(str(header))] + [len(str(r[index - 1])) for r in table.rows[:200] if index - 1 < len(r)])
            sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def csv_safe(text: str) -> str:
    """Excel runs a CSV cell that starts with =, @ (or + / - that isn't a plain number) as a formula
    ("=cmd|..." can start programs). A leading apostrophe makes it text."""
    if not text:
        return text
    first = text[0]
    if first in ("=", "@", chr(9), chr(13)) or (first in "+-" and not _NUMBER.match(text.strip())):
        return "'" + text
    return text


def _write_csv(table: Table, path: Path) -> None:
    # utf-8 with a BOM: Excel then opens accented and non-Latin text correctly
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([csv_safe(h) for h in table.headers])
        writer.writerows([[csv_safe(c) for c in row] for row in table.rows])


def write_csv(tables: list[Table], path: Path) -> Path:
    """One table -> `path` as a CSV. Several -> `path` with a .zip suffix holding one CSV per table.
    Returns the file written."""
    if not tables:
        raise NoTablesError("there is no table to put in a CSV file")
    path.parent.mkdir(parents=True, exist_ok=True)
    if len(tables) == 1:
        _write_csv(tables[0], path)
        return path
    archive = path.with_suffix(".zip")
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for table, name in zip(tables, _sheet_names(tables)):
            member = path.parent / f"{name}.csv"
            _write_csv(table, member)
            bundle.write(member, arcname=f"{re.sub(r'[^A-Za-z0-9_. -]+', '-', name).strip() or 'table'}.csv")
            member.unlink()
    return archive


def export_tables(markdown: str, out_dir: Path, stem: str, fmt: str) -> Path:
    """The tables of `markdown` as `<stem>.xlsx` or `<stem>.csv` (a .zip when there are several)."""
    tables = extract_tables(markdown)
    if not tables:
        raise NoTablesError("this answer has no table to export")
    if fmt == "xlsx":
        return write_xlsx(tables, out_dir / f"{stem}.xlsx")
    if fmt == "csv":
        return write_csv(tables, out_dir / f"{stem}.csv")
    raise ValueError(f"unsupported table format: {fmt!r}")
